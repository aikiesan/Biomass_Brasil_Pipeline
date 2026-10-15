"""Municipal sanitation surveys -> one long table on harmonised variables.

The federal sanitation survey changed system after reference year 2022:

* **SNIS** (to 2022): one "Consolidado por Município" CSV per year, UTF-16-LE,
  ';'-delimited, Brazilian number format ("8.000,00"). Columns are headed
  "<code> - <label>" and selected by code, so columns added or dropped between
  years are harmless. The municipality key is the 6-digit "Código do Município"
  (the 7-digit "Código do IBGE" column is empty before ~2019); it is resolved to
  the 7-digit geocode through the mesh, whose 6-digit prefixes are unique. Each
  file ends with a "TOTAL da AMOSTRA" row: SNIS's own sum of the municipal rows,
  kept to check the parse and never ingested as data.
* **SINISA** (from 2023): XLSX workbooks with a title block, a row of field codes
  ("Cod_IBGE", "GTR1028", ...; codes carry a trailing "*" for mandatory fields
  and stray spaces) and one row per municipality. Municipalities that did not
  answer a module carry the literal "null". The consolidated UF / macro-region /
  Brazil workbooks are SINISA's own sums, used to check the municipal ones.

Reporting is voluntary in both surveys, so a blank cell means "not reported",
never zero: only reported cells become rows. Which field of each survey feeds
which harmonised variable is configured in config/sanitation.yaml.
"""

from __future__ import annotations

import csv
from collections.abc import Iterable, Sequence
from pathlib import Path

import numpy as np
import pandas as pd

SNIS_ENCODING = "utf-16-le"
SNIS_KEY = "Código do Município"
SINISA_KEY = "cod_ibge"
SINISA_RESPONDED = "Sim/Não"
NOT_REPORTED = {"", "null", "none", "nan"}

COLUMNS = ["geocode", "year", "survey", "module", "field", "value", "source_file"]


# ---------------------------------------------------------------------------
# cells
# ---------------------------------------------------------------------------
def clean_code(cell) -> str:
    """'GTR1028* ' -> 'GTR1028'; 'CO119 - Quantidade total ...' -> 'CO119'."""
    text = str(cell if cell is not None else "").replace("\xa0", " ").strip()
    text = text.split(" - ")[0].strip()
    return text.rstrip("*").strip()


def parse_number(raw, brazilian: bool = False) -> tuple[float, str]:
    """(value, flag). Blank and SINISA's "null" are not reported (NaN), never zero.

    `brazilian=True` reads text as "1.234,5" — SNIS exports always are, and there a
    single dot is a thousands separator ("203.080" is 203,080, not 203.08). SINISA
    cells are numbers, or text with a decimal point.
    """
    if raw is None:
        return np.nan, "not_reported"
    if isinstance(raw, bool):
        return np.nan, f"unparsed:{raw}"
    if isinstance(raw, int | float):
        return (np.nan, "not_reported") if np.isnan(raw) else (float(raw), "ok")
    text = str(raw).replace("\xa0", " ").strip()
    if text.lower() in NOT_REPORTED:
        return np.nan, "not_reported"
    if brazilian or "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text), "ok"
    except ValueError:
        return np.nan, f"unparsed:{raw}"


def _geocode(cell) -> int | None:
    text = str(cell if cell is not None else "").strip()
    if text.endswith(".0"):
        text = text[:-2]
    return int(text) if len(text) == 7 and text.isdigit() else None


# ---------------------------------------------------------------------------
# SNIS
# ---------------------------------------------------------------------------
def parse_snis_rows(rows: Iterable[Sequence[str]], year: int, fields: Sequence[str]) -> tuple[pd.DataFrame, dict]:
    """Rows of a SNIS export (header first) -> (long frame keyed on the 6-digit code, metadata).

    Metadata: `totals` ({field: value} of the "TOTAL da AMOSTRA" row), `summary_rows`,
    `missing_fields` (requested codes absent from this year's header), `unparsed`.
    """
    rows = iter(rows)
    header = [clean_code(h) for h in next(rows)]
    index = {code: i for i, code in reversed(list(enumerate(header)))}  # first occurrence wins
    if SNIS_KEY not in index:
        raise ValueError(f"SNIS export without a '{SNIS_KEY}' column; header starts {header[:5]}")
    key_i = index[SNIS_KEY]
    cols = {f: index[f] for f in fields if f in index}
    records, totals, summary_rows, unparsed = [], {}, 0, []
    for row in rows:
        if not any(c.strip() for c in row):
            continue
        key = row[key_i].strip() if key_i < len(row) else ""
        cells = {f: parse_number(row[i] if i < len(row) else None, brazilian=True) for f, i in cols.items()}
        if not (len(key) == 6 and key.isdigit()):
            summary_rows += 1
            totals = {f: v for f, (v, flag) in cells.items() if flag == "ok"}
            continue
        for f, (value, flag) in cells.items():
            if flag == "not_reported":
                continue
            if flag != "ok":
                unparsed.append([int(key), f, flag])
                continue
            records.append((int(key), f, value))
    df = pd.DataFrame.from_records(records, columns=["code6", "field", "value"])
    df["year"] = year
    meta = {"totals": totals, "summary_rows": summary_rows, "unparsed": unparsed[:50], "n_unparsed": len(unparsed),
            "missing_fields": [f for f in fields if f not in index]}
    return df, meta


def read_snis(path: Path, year: int, fields: Sequence[str]) -> tuple[pd.DataFrame, dict]:
    with open(path, encoding=SNIS_ENCODING, newline="") as fh:
        df, meta = parse_snis_rows(csv.reader(fh, delimiter=";"), year, fields)
    df["source_file"] = Path(path).name
    return df, meta


def code6_index(geocodes: Iterable[int]) -> dict[int, int]:
    """{6-digit prefix: 7-digit geocode}. Raises if two geocodes share a prefix."""
    out: dict[int, int] = {}
    for g in geocodes:
        p = int(g) // 10
        if p in out:
            raise ValueError(f"geocodes {out[p]} and {g} share the 6-digit prefix {p}")
        out[p] = int(g)
    return out


# ---------------------------------------------------------------------------
# SINISA
# ---------------------------------------------------------------------------
def parse_sinisa_rows(rows: Iterable[Sequence], year: int, fields: Sequence[str]) -> tuple[pd.DataFrame, dict]:
    """Rows of a SINISA municipal sheet -> (long frame, metadata).

    The code row is the first row with a "Cod_IBGE" cell; everything above it is the
    title block. Metadata: `rows` (municipality rows), `respondents` (rows answering
    "Sim" in the module column, None when the sheet has no such column),
    `duplicates`, `missing_fields`, `unparsed`.
    """
    rows = iter(rows)
    header = None
    for row in rows:
        codes = [clean_code(c) for c in row]
        if any(c.lower() == SINISA_KEY for c in codes):
            header = codes
            break
    if header is None:
        raise ValueError("no 'Cod_IBGE' code row found")
    index = {code: i for i, code in reversed(list(enumerate(header)))}
    key_i = next(i for i, c in enumerate(header) if c.lower() == SINISA_KEY)
    resp_i = index.get(SINISA_RESPONDED)
    cols = {f: index[f] for f in fields if f in index}
    records, keys, respondents, unparsed = [], [], 0, []
    for row in rows:
        geo = _geocode(row[key_i] if key_i < len(row) else None)
        if geo is None:
            continue  # blank lines, notes
        keys.append(geo)
        if resp_i is not None and str(row[resp_i] or "").strip().lower() == "sim":
            respondents += 1
        for f, i in cols.items():
            value, flag = parse_number(row[i] if i < len(row) else None)
            if flag == "not_reported":
                continue
            if flag != "ok":
                unparsed.append([geo, f, flag])
                continue
            records.append((geo, f, value))
    df = pd.DataFrame.from_records(records, columns=["geocode", "field", "value"])
    df["year"] = year
    meta = {"rows": len(keys), "respondents": respondents if resp_i is not None else None,
            "duplicates": len(keys) - len(set(keys)), "n_unparsed": len(unparsed), "unparsed": unparsed[:50],
            "missing_fields": [f for f in fields if f not in index]}
    return df, meta


def _sheet_rows(path: Path, sheet: str) -> list[tuple]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"{Path(path).name}: no sheet {sheet!r} (has {wb.sheetnames})")
        return list(wb[sheet].iter_rows(values_only=True))
    finally:
        wb.close()


def read_sinisa(path: Path, sheet: str, year: int, fields: Sequence[str]) -> tuple[pd.DataFrame, dict]:
    df, meta = parse_sinisa_rows(_sheet_rows(path, sheet), year, fields)
    df["source_file"] = Path(path).name
    return df, meta


MACRO_REGIONS = {"Norte": "N", "Nordeste": "NE", "Sudeste": "SE", "Sul": "S", "Centro-Oeste": "CO"}


def parse_sinisa_published(rows: Iterable[Sequence], fields: Sequence[str]) -> dict[str, dict[str, float]]:
    """SINISA consolidated sheet -> {unit: {field: value}}: "BR", UF abbreviations ("SP"), and
    macro-regions as "R:<abbr>" ("R:SE").

    The code row is the first row holding one of `fields`; the unit name is the first
    column and its abbreviation the second. Macro-regions are told apart by name,
    because their abbreviations collide with UFs (Sudeste and Sergipe are both "SE").
    """
    rows = iter(rows)
    header = None
    for row in rows:
        codes = [clean_code(c) for c in row]
        if any(c in fields for c in codes):
            header = codes
            break
    if header is None:
        raise ValueError(f"none of {list(fields)} found in the consolidated sheet")
    cols = {f: header.index(f) for f in fields if f in header}
    out: dict[str, dict[str, float]] = {}
    for row in rows:
        name = str(row[0] if row and row[0] is not None else "").strip()
        sigla = str(row[1] if len(row) > 1 and row[1] is not None else "").strip()
        if not (1 <= len(sigla) <= 2 and sigla.isalpha() and sigla.isupper()):
            continue
        if name in MACRO_REGIONS:
            sigla = f"R:{MACRO_REGIONS[name]}"
        values = {}
        for f, i in cols.items():
            value, flag = parse_number(row[i] if i < len(row) else None)
            if flag == "ok":
                values[f] = value
        out[sigla] = values
    return out


def read_sinisa_published(path: Path, sheet: str, fields: Sequence[str]) -> dict[str, dict[str, float]]:
    return parse_sinisa_published(_sheet_rows(path, sheet), fields)


# ---------------------------------------------------------------------------
# harmonisation
# ---------------------------------------------------------------------------
def crosswalk(variables: dict) -> pd.DataFrame:
    """config/sanitation.yaml `variables` -> one row per (survey, module, field, variable)."""
    rows = []
    for var, spec in variables.items():
        common = {"variable": var, "unit": spec["unit"], "match": spec["match"]}
        if spec.get("snis"):
            rows.append({"survey": "SNIS", "module": "snis", "field": spec["snis"], **common})
        if spec.get("sinisa"):
            rows.append({"survey": "SINISA", "module": spec["sinisa"]["module"], "field": spec["sinisa"]["field"],
                         **common})
    xw = pd.DataFrame(rows)
    dup = xw[xw.duplicated(["survey", "module", "field"], keep=False)]
    if len(dup):
        raise ValueError(f"a survey field feeds more than one variable: {dup.to_dict('records')}")
    return xw


def fields_for(xw: pd.DataFrame, survey: str, module: str, extra: Iterable[str] = ()) -> list[str]:
    sel = xw[(xw["survey"] == survey) & (xw["module"] == module)]["field"]
    return sorted(set(sel) | set(extra))


def harmonise(long: pd.DataFrame, xw: pd.DataFrame) -> pd.DataFrame:
    """Attach the harmonised variable to every survey cell; cells of fields not in the crosswalk
    (read only for identity checks) are dropped."""
    out = long.merge(xw, on=["survey", "module", "field"], how="inner")
    cols = ["geocode", "year", "survey", "variable", "unit", "value", "field", "match", "source_file"]
    out = out[cols].sort_values(["variable", "geocode", "year"]).reset_index(drop=True)
    out["geocode"], out["year"] = out["geocode"].astype("int64"), out["year"].astype("int16")
    return out


def identity_breaks(long: pd.DataFrame, total: str, parts: Sequence[str], rel: float = 0.01) -> pd.DataFrame:
    """(geocode, year) where all of total and parts are reported and |total − Σparts| > rel × Σparts."""
    w = long[long["field"].isin([total, *parts])].pivot_table(index=["geocode", "year"], columns="field",
                                                             values="value", aggfunc="first")
    if not {total, *parts} <= set(w.columns):
        return w.iloc[0:0]
    w = w[[total, *parts]].dropna()
    s = w[list(parts)].sum(axis=1)
    off = (w[total] - s).abs() > rel * s.clip(lower=1)
    # filter `s` too: assigning a full-length Series to an EMPTY frame would adopt its index
    return w[off].assign(parts_sum=s[off])


def sums_by_uf(long: pd.DataFrame, uf_of: pd.Series, fields: Sequence[str],
               region_of: pd.Series | None = None) -> dict[str, dict[str, float]]:
    """{uf | 'R:<region>' | 'BR': {field: municipal sum}} — keyed like parse_sinisa_published."""
    sel = long[long["field"].isin(fields)]
    out: dict[str, dict[str, float]] = {}
    keys = [("", uf_of)] + ([("R:", region_of)] if region_of is not None else [])
    for prefix, unit_of in keys:
        by = sel.assign(unit=sel["geocode"].map(unit_of)).groupby(["unit", "field"])["value"].sum().unstack("field")
        out.update({f"{prefix}{u}": {f: float(v) for f, v in r.items() if pd.notna(v)} for u, r in by.iterrows()})
    out["BR"] = {f: float(v) for f, v in sel.groupby("field")["value"].sum().items()}
    return out


def continuity(h: pd.DataFrame) -> pd.DataFrame:
    """Series break: the last SNIS year against the first SINISA year, per variable, over the
    municipalities that report a positive value in both."""
    rows = []
    for var, g in h.groupby("variable", observed=True):
        surveys = set(g["survey"])
        if surveys != {"SNIS", "SINISA"}:
            continue
        y0 = int(g.loc[g["survey"] == "SNIS", "year"].max())
        y1 = int(g.loc[g["survey"] == "SINISA", "year"].min())
        a = g[(g["survey"] == "SNIS") & (g["year"] == y0)].set_index("geocode")["value"]
        b = g[(g["survey"] == "SINISA") & (g["year"] == y1)].set_index("geocode")["value"]
        both = pd.concat([a.rename("snis"), b.rename("sinisa")], axis=1, join="inner")
        both = both[(both["snis"] > 0) & (both["sinisa"] > 0)]
        r = both["sinisa"] / both["snis"]
        rows.append({
            "variable": var, "match": g["match"].iloc[0], "snis_year": y0, "sinisa_year": y1,
            "n_snis": int(len(a)), "n_sinisa": int(len(b)), "n_both": int(len(both)),
            "snis_total_both": float(both["snis"].sum()), "sinisa_total_both": float(both["sinisa"].sum()),
            "ratio_of_totals": float(both["sinisa"].sum() / both["snis"].sum()) if len(both) else np.nan,
            "median_ratio": float(r.median()) if len(r) else np.nan,
            "p10_ratio": float(r.quantile(0.10)) if len(r) else np.nan,
            "p90_ratio": float(r.quantile(0.90)) if len(r) else np.nan,
            "share_within_25pct": float(((r - 1).abs() <= 0.25).mean()) if len(r) else np.nan,
        })
    return pd.DataFrame(rows)
