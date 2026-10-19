"""MapBiomas municipal coverage statistics (Collection 11 XLSX) -> long table.

The workbook sheet COVERAGE_11 has one row per municipality x biome x class and
one column per year (y1985 ... y2025), values in hectares. We stream it with
openpyxl in read-only mode (the sheet XML is ~170 MB) and keep only the
configured years. Areas are summed over biomes later; the biome is kept so
that biome can be used as an explanatory variable.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

import numpy as np
import pandas as pd

log = logging.getLogger("bbp")

YEAR_COL = re.compile(r"^y?(\d{4})$")
REQUIRED = {"geocode", "biome", "class"}


def _header_index(header: list) -> tuple[dict[str, int], dict[int, int]]:
    names = [str(h).strip() if h is not None else "" for h in header]
    idx = {n: i for i, n in enumerate(names)}
    missing = REQUIRED - set(idx)
    if missing:
        raise ValueError(f"MapBiomas sheet lacks columns {sorted(missing)}; header = {names}")
    years = {}
    for i, n in enumerate(names):
        m = YEAR_COL.match(n)
        if m:
            years[int(m.group(1))] = i
    return idx, years


def parse_coverage_rows(rows, years_wanted: list[int]) -> pd.DataFrame:
    """Parse an iterator of rows (first row = header)."""
    rows = iter(rows)
    header = list(next(rows))
    idx, year_cols = _header_index(header)
    use_years = [y for y in years_wanted if y in year_cols]
    if not use_years:
        raise ValueError(f"None of the years {years_wanted[0]}–{years_wanted[-1]} are in the sheet")
    gi, bi, ci = idx["geocode"], idx["biome"], idx["class"]
    li = idx.get("class_level_4")
    ycols = [year_cols[y] for y in use_years]

    geocodes, biomes, classes, values = [], [], [], []
    labels: dict[int, str] = {}
    for n, row in enumerate(rows, start=1):
        g = row[gi]
        if g is None or str(g).strip() == "":
            continue
        geocodes.append(int(float(g)))
        biomes.append(row[bi])
        cls = int(float(row[ci]))
        classes.append(cls)
        if li is not None and cls not in labels:
            labels[cls] = str(row[li])
        values.append([np.nan if row[c] is None else float(row[c]) for c in ycols])
        if n % 50000 == 0:
            log.info("  MapBiomas rows read: %s", f"{n:,}")

    wide = pd.DataFrame(np.asarray(values, dtype="float64"), columns=use_years)
    wide["geocode"] = np.asarray(geocodes, dtype="int64")
    wide["biome"] = biomes
    wide["class_id"] = np.asarray(classes, dtype="int16")
    long = wide.melt(id_vars=["geocode", "biome", "class_id"], var_name="year", value_name="area_ha")
    long["year"] = long["year"].astype("int16")
    long = long.dropna(subset=["area_ha"])
    long["biome"] = long["biome"].astype("category")
    long["class_label"] = long["class_id"].map(labels).astype("category")
    return long.reset_index(drop=True)


def read_coverage_xlsx(path: Path, sheet: str, years: list[int]) -> pd.DataFrame:
    import openpyxl

    log.info("Reading %s [%s] (streaming; this takes a few minutes)", path.name, sheet)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"Sheet {sheet!r} not in {path.name}: {wb.sheetnames}")
        return parse_coverage_rows(wb[sheet].iter_rows(values_only=True), years)
    finally:
        wb.close()


IRRIGATION_REQUIRED = {"MUNICIPALITY_GEOCODE", "BIOME", "IRRIGATION_CLASS_NAME", "AGRICULTURE_CLASS_ID",
                       "AGRICULTURE_CLASS_NAME", "COLLECTION"}


def parse_irrigation_rows(rows, years_wanted: list[int]) -> pd.DataFrame:
    """MapBiomas irrigation-systems statistics (sheet IRRIGATION_SYSTEMS) -> long table.

    One row per municipality x biome x irrigation system x irrigated crop, one column per year
    (hectares). Returns geocode, biome, irrigation_class, agriculture_class_id, agriculture_class,
    collection, year, area_ha.
    """
    rows = iter(rows)
    names = [str(h).strip() if h is not None else "" for h in next(rows)]
    idx = {n: i for i, n in enumerate(names)}
    missing = IRRIGATION_REQUIRED - set(idx)
    if missing:
        raise ValueError(f"MapBiomas irrigation sheet lacks columns {sorted(missing)}; header = {names}")
    year_cols = {int(m.group(1)): i for i, n in enumerate(names) if (m := YEAR_COL.match(n))}
    use_years = [y for y in years_wanted if y in year_cols]
    if not use_years:
        raise ValueError(f"None of the years {years_wanted[0]}–{years_wanted[-1]} are in the irrigation sheet")
    keys, values = [], []
    for row in rows:
        g = row[idx["MUNICIPALITY_GEOCODE"]]
        if g is None or str(g).strip() == "":
            continue
        keys.append((int(float(g)), row[idx["BIOME"]], row[idx["IRRIGATION_CLASS_NAME"]],
                     int(float(row[idx["AGRICULTURE_CLASS_ID"]])), row[idx["AGRICULTURE_CLASS_NAME"]],
                     int(float(row[idx["COLLECTION"]]))))
        values.append([np.nan if row[year_cols[y]] is None else float(row[year_cols[y]]) for y in use_years])
    wide = pd.DataFrame(keys, columns=["geocode", "biome", "irrigation_class", "agriculture_class_id",
                                       "agriculture_class", "collection"])
    wide[use_years] = np.asarray(values, dtype="float64").reshape(len(keys), len(use_years))
    long = wide.melt(id_vars=list(wide.columns[:6]), var_name="year", value_name="area_ha").dropna(subset=["area_ha"])
    long["year"] = long["year"].astype("int16")
    long["geocode"] = long["geocode"].astype("int64")
    for c in ("biome", "irrigation_class", "agriculture_class"):
        long[c] = long[c].astype("category")
    return long.reset_index(drop=True)


def read_irrigation_xlsx(path: Path, years: list[int], sheet: str = "IRRIGATION_SYSTEMS") -> pd.DataFrame:
    import openpyxl

    log.info("Reading %s [%s]", path.name, sheet)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"Sheet {sheet!r} not in {path.name}: {wb.sheetnames}")
        return parse_irrigation_rows(wb[sheet].iter_rows(values_only=True), years)
    finally:
        wb.close()


def read_legend(path: Path) -> pd.DataFrame:
    leg = pd.read_csv(path)
    leg = leg.rename(columns={"class_name_en": "class_name"})
    leg["class_id"] = leg["class_id"].astype("int16")
    return leg
